"""Permission checked planning operations and same-document amendments."""
import csv
import io
import json

import frappe
from frappe.utils import cint, get_datetime, now_datetime

from sales_performance.services.target_editor import (
    HEADERS, TABLES, DIMENSIONS, APPROVAL_TOKEN, apply_snapshot, snapshot, normalize, transfer, number,
)
from sales_performance.services.planning_engine import make_row_key
from sales_performance.services.numbers import nflt as flt


def plan(name, write=True, lock=False):
    if lock:
        frappe.db.sql("select name from `tabSales Target Planning` where name=%s for update", name)
    doc = frappe.get_doc("Sales Target Planning", name)
    doc.check_permission("write" if write else "read")
    return doc


def approver():
    if frappe.session.user != "Administrator" and not set(frappe.get_roles()).intersection(
        {"Sales Target Approver", "Sales Performance Admin", "System Manager"}
    ):
        frappe.throw("Target approver access is required", frappe.PermissionError)


def target_amender():
    if frappe.session.user != "Administrator" and not set(frappe.get_roles()).intersection(
        {"Sales Target Editor", "Sales Target Manager", "Sales Performance Admin", "System Manager"}
    ):
        frappe.throw("Sales Target Manager access is required to amend approved targets", frappe.PermissionError)


def edited(doc, payload, baseline=None, effective_month=1):
    data = frappe.parse_json(payload) if isinstance(payload, str) else payload
    if not isinstance(data, dict):
        frappe.throw("Invalid target data")
    apply_snapshot(doc, data)
    if baseline and (doc.entry_basis == "Monthly" or doc.distribution_method == "Manual Monthly"):
        _apply_annual_amendments_to_months(doc, baseline, effective_month)
    normalize(doc, pricing_start_month=effective_month)
    doc.check_permission("write")
    doc._action = "save"
    doc._validate_links()
    status = doc.status
    doc.status = "Draft"
    doc._validate_business_rules()
    doc.status = status
    return doc


def _apply_annual_amendments_to_months(doc, baseline, effective_month):
    """Keep month shares when an editor changes an annual Proposal Details value."""
    from collections import defaultdict
    from sales_performance.services.precision import get_currency_precision, get_qty_precision

    old_rows = {row.get("row_key"): row for row in baseline.get("proposal_details", []) if row.get("row_key")}
    months = defaultdict(list)
    for month in doc.monthly_details or []:
        months[month.row_key].append(month)
    amount_entry = doc.value_basis == "Quantity and Amount"
    start = int(effective_month or 1)
    for row in doc.proposal_details or []:
        old = old_rows.get(row.row_key)
        if not old:
            continue
        row_months = months.get(row.row_key, [])
        if not row_months:
            continue
        if flt(row.approved_target_qty) != flt(old.get("approved_target_qty")):
            _scale_months_to_total(row_months, "target_qty", flt(row.approved_target_qty), start,
                                   get_qty_precision(row.uom))
        if amount_entry:
            if flt(row.approved_target_amount) != flt(old.get("approved_target_amount")):
                _scale_months_to_total(row_months, "target_amount", flt(row.approved_target_amount), start,
                                       get_currency_precision())
        else:
            new_rate = flt(row.approved_target_rate)
            if new_rate != flt(old.get("approved_target_rate")):
                for month in row_months:
                    if int(month.month_number) >= start:
                        month.target_rate = new_rate
                        month.target_amount = flt(month.target_qty * new_rate, get_currency_precision())


def _scale_months_to_total(months, field, annual_total, start_month, precision):
    protected = sum(flt(month.get(field)) for month in months if int(month.month_number) < start_month)
    remaining = annual_total - protected
    if remaining < -(10 ** -max(precision, 1)):
        frappe.throw("The annual target is below the amount already protected before the effective month.")
    remaining = max(remaining, 0)
    active = sorted((month for month in months if int(month.month_number) >= start_month), key=lambda month: int(month.month_number))
    if not active:
        return
    baseline = sum(flt(month.get(field)) for month in active)
    from sales_performance.services.distribution_engine import reconcile_to_total

    values = [remaining / len(active) if not baseline else remaining * flt(month.get(field)) / baseline
              for month in active]
    for month, value in zip(active, reconcile_to_total(values, remaining, precision)):
        month.update({field: value})


def amendment_doc(name, amendment, modified=None):
    doc = plan(name, lock=True)
    record = frappe.get_doc("Sales Target Amendment", amendment)
    if record.planning != name or record.status not in ("Draft", "Under Review"):
        frappe.throw("This amendment is no longer editable")
    if doc.status != "Approved":
        frappe.throw("The plan is no longer approved")
    if get_datetime(record.base_modified) != get_datetime(doc.modified):
        frappe.throw("The approved plan changed. Discard this amendment and start again.")
    if modified and get_datetime(record.modified) != get_datetime(modified):
        frappe.throw("Another user changed this amendment. Reload it before continuing.")
    return doc, record


@frappe.whitelist()
def start_amendment(name, reason=None, effective_month=None):
    target_amender()
    doc = plan(name, lock=True)
    if doc.status != "Approved":
        frappe.throw("Only approved targets need an amendment")
    existing = frappe.db.get_value("Sales Target Amendment", {"planning": name, "status": ["in", ["Draft", "Under Review"]]}, "name")
    if existing:
        record = frappe.get_doc("Sales Target Amendment", existing)
    else:
        if not reason or not str(reason).strip():
            frappe.throw("An amendment reason is required")
        month = number(effective_month, "Effective month")
        if month != int(month) or not 1 <= month <= 12:
            frappe.throw("Select an effective month from 1–12")
        data = json.dumps(snapshot(doc), default=str)
        record = frappe.get_doc(dict(doctype="Sales Target Amendment", planning=name, status="Draft",
            reason=reason, effective_month=int(month), base_modified=str(doc.modified),
            before_values=data, proposed_values=data))
        record.insert(ignore_permissions=True)
    return amendment_result(record, doc)


def amendment_result(record, fallback_doc=None):
    payload = json.loads(record.proposed_values)
    if fallback_doc:
        from sales_performance.services.target_editor import DISPLAY_FIELDS
        live_rows = {row.row_key: row for row in fallback_doc.proposal_details or [] if row.row_key}
        for row in payload.get("proposal_details", []):
            live = live_rows.get(row.get("row_key"))
            if live:
                for field in DISPLAY_FIELDS:
                    if row.get(field) is None:
                        row[field] = live.get(field)
    return dict(name=record.name, planning=record.planning, modified=str(record.modified), status=record.status, reason=record.reason,
                effective_month=record.effective_month, payload=payload)


@frappe.whitelist()
def pending_amendment(name):
    """Show saved changes explicitly without replacing the approved reporting data."""
    doc = plan(name, write=False)
    if doc.status != "Approved":
        return None
    pending = frappe.db.get_value("Sales Target Amendment",
        {"planning": name, "status": ["in", ["Draft", "Under Review"]]},
        ["name", "status", "reason", "effective_month"], as_dict=True)
    return pending


def target_map(data):
    return {(tuple(r.get(f) or "" for f in DIMENSIONS), cint(r.get("month_number"))):
            (float(r.get("target_qty") or 0), float(r.get("target_amount") or 0))
            for r in data.get("monthly_details", [])}


def validate_effective_month(before, after, effective_month):
    old, new = target_map(before), target_map(after)
    for key in old.keys() | new.keys():
        if key[1] < effective_month and old.get(key, (0, 0)) != new.get(key, (0, 0)):
            frappe.throw("Targets before the effective month cannot change. Start an Entire Year amendment to correct earlier months.")


@frappe.whitelist()
def save_amendment(name, amendment, modified, payload, submit=False):
    target_amender()
    doc, record = amendment_doc(name, amendment, modified)
    edited(doc, payload, effective_month=record.effective_month)
    for row in doc.proposal_details:
        if row.override_percent and not row.override_reason:
            row.override_reason = record.reason
    data = snapshot(doc)
    validate_effective_month(json.loads(record.before_values), data, record.effective_month)
    record.proposed_values = json.dumps(data, default=str)
    record.status = "Under Review" if cint(submit) else "Draft"
    record.save(ignore_permissions=True)
    return amendment_result(record)


@frappe.whitelist()
def approve_amendment(name, amendment, modified):
    approver()
    doc, record = amendment_doc(name, amendment, modified)
    before = json.loads(record.before_values)
    edited(doc, record.proposed_values, effective_month=record.effective_month)
    validate_effective_month(before, snapshot(doc), record.effective_month)
    doc.flags.authorized_target_action = APPROVAL_TOKEN
    doc.amendment_count = cint(doc.amendment_count) + 1
    doc.approved_by, doc.approved_on = frappe.session.user, now_datetime()
    try:
        # Version would serialize thousands of monthly child rows into one
        # oversized SQL packet. The amendment record already retains the full
        # before/after snapshot and approval history.
        doc.save(ignore_version=True)
    finally:
        doc.flags.authorized_target_action = False
    from sales_performance.services.target_sync import sync_official_targets
    sync_official_targets(doc)
    record.status = "Applied"
    record.amendment_number = doc.amendment_count
    record.approved_by, record.approved_on = doc.approved_by, doc.approved_on
    record.proposed_values = json.dumps(snapshot(doc), default=str)
    record.save(ignore_permissions=True)
    return doc.name


@frappe.whitelist()
def apply_amendment(name, amendment, modified, payload):
	"""Save and apply an approved-plan amendment in one role-checked action."""
	target_amender()
	doc, record = amendment_doc(name, amendment, modified)
	before = json.loads(record.before_values)
	edited(doc, payload, baseline=before, effective_month=record.effective_month)
	validate_effective_month(before, snapshot(doc), record.effective_month)
	doc.flags.authorized_target_action = APPROVAL_TOKEN
	doc.amendment_count = cint(doc.amendment_count) + 1
	doc.approved_by, doc.approved_on = frappe.session.user, now_datetime()
	try:
		doc.save(ignore_version=True)
	finally:
		doc.flags.authorized_target_action = False
	from sales_performance.services.target_sync import sync_official_targets
	sync_official_targets(doc)
	record.status = "Applied"
	record.amendment_number = doc.amendment_count
	record.approved_by, record.approved_on = doc.approved_by, doc.approved_on
	record.proposed_values = json.dumps(snapshot(doc), default=str)
	record.save(ignore_permissions=True)
	return doc.name


@frappe.whitelist()
def discard_amendment(name, amendment, modified):
    target_amender()
    doc = plan(name, lock=True)
    record = frappe.get_doc("Sales Target Amendment", amendment)
    if record.planning != doc.name or record.status not in ("Draft", "Under Review"):
        frappe.throw("This amendment cannot be discarded")
    if get_datetime(record.modified) != get_datetime(modified):
        frappe.throw("Reload the amendment before discarding")
    record.status = "Discarded"
    record.save(ignore_permissions=True)


@frappe.whitelist()
def history(name):
    plan(name, write=False)
    return frappe.get_all("Sales Target Amendment", filters={"planning": name},
        fields=["name", "status", "reason", "effective_month", "amendment_number", "owner", "creation", "approved_by", "approved_on"],
        order_by="creation desc")


@frappe.whitelist()
def amendment_preview(name, amendment):
    plan(name, write=False)
    record = frappe.get_doc("Sales Target Amendment", amendment)
    if record.planning != name:
        frappe.throw("Invalid amendment", frappe.PermissionError)
    before, after = json.loads(record.before_values), json.loads(record.proposed_values)
    old, new = target_map(before), target_map(after)
    changes = []
    for key in sorted(old.keys() | new.keys()):
        if old.get(key, (0, 0)) != new.get(key, (0, 0)):
            changes.append(dict(zip(DIMENSIONS, key[0]), month=key[1],
                before_qty=old.get(key, (0, 0))[0], after_qty=new.get(key, (0, 0))[0],
                before_amount=old.get(key, (0, 0))[1], after_amount=new.get(key, (0, 0))[1]))
    headers = [{"field": f, "before": before.get(f), "after": after.get(f)} for f in HEADERS if before.get(f) != after.get(f)]
    row_changes = []
    for table in ("proposal_details", "growth_rules", "custom_percents"):
        identity = "row_key" if table == "proposal_details" else "item_group" if table == "growth_rules" else "month"
        old_rows = {r.get(identity): r for r in before.get(table, [])}
        new_rows = {r.get(identity): r for r in after.get(table, [])}
        for key in old_rows.keys() | new_rows.keys():
            old_row, new_row = old_rows.get(key, {}), new_rows.get(key, {})
            for field in TABLES[table]:
                if old_row.get(field) != new_row.get(field):
                    row_changes.append({"table": table, "row": new_row.get("item_code") or old_row.get("item_code") or key,
                        "sales_person": new_row.get("sales_person") or old_row.get("sales_person"), "field": field,
                        "before": old_row.get(field), "after": new_row.get(field)})
    return {"changes": changes, "settings": headers, "row_changes": row_changes, "reason": record.reason,
            "notice": "Submitted and paid incentive payouts are unchanged. Historical corrections may require a separate payout adjustment."}


@frappe.whitelist()
def preview_edit(name, payload, operation="normalize", options=None):
    doc = plan(name)
    if doc.status in ("Cancelled", "Superseded"):
        frappe.throw("This plan cannot be edited")
    options = frappe.parse_json(options) or {}
    data = frappe.parse_json(payload)
    if doc.status != "Approved":
        for field in ("company", "fiscal_year"):
            if data.get(field):
                doc.set(field, data[field])
    apply_snapshot(doc, data)
    result = {}
    if operation == "transfer":
        if not options.get("reason"):
            frappe.throw("A transfer reason is required")
        if options.get("scope") == "Selected Rows" and not options.get("row_keys"):
            frappe.throw("Select proposal rows first, or choose All Matching Rows.")
        result = transfer(doc, options.get("source"), options.get("destination"), options.get("start_month"),
                          options.get("percent", 100), options.get("row_keys"))
        for row in doc.proposal_details:
            if row.sales_person in (options.get("source"), options.get("destination")):
                row.override_reason = options["reason"]
    elif operation == "distribute":
        normalize(doc, force_distribution=True)
    elif operation == "copy":
        source = plan(options.get("source"), write=False)
        if source.company != doc.company:
            frappe.throw("Copy targets only between plans for the same company")
        data = snapshot(source)
        for f in ("title", "sales_person", "territory", "item_group", "customer_group", "notes"):
            data.pop(f, None)
        apply_snapshot(doc, data)
        doc.target_source = "Copy Existing Plan"
        normalize(doc)
    elif operation == "import":
        result = import_rows(doc, options)
    elif operation == "normalize":
        normalize(doc)
    else:
        frappe.throw("Unknown editing operation")
    doc.check_permission("write")
    doc._action = "save"
    doc._validate_links()
    return {"payload": snapshot(doc), "summary": result}


def read_import(name, file_url):
    file_name = frappe.db.get_value("File", {"file_url": file_url, "attached_to_doctype": "Sales Target Planning", "attached_to_name": name}, "name")
    if not file_name:
        frappe.throw("Upload the import file as an attachment to this plan")
    file = frappe.get_doc("File", file_name)
    file.check_permission("read")
    content = file.get_content()
    if file.file_name.lower().endswith(".xlsx"):
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            values = list(book.active.iter_rows(values_only=True))
        finally:
            book.close()
    elif file.file_name.lower().endswith(".csv"):
        text = content.decode("utf-8-sig") if isinstance(content, bytes) else content.lstrip("\ufeff")
        values = list(csv.reader(io.StringIO(text)))
    else:
        frappe.throw("Use an .xlsx or UTF-8 .csv file")
    if len(values) < 2 or len(values) > 10001:
        frappe.throw("Import must contain 1–10,000 data rows")
    headers = [str(v or "").strip().lower().replace(" ", "_") for v in values[0]]
    if len(headers) != len(set(headers)):
        frappe.throw("Duplicate import column names")
    return [dict(zip(headers, values)) for values in values[1:] if any(v not in (None, "") for v in values)]


def import_rows(doc, options):
    rows = read_import(doc.name, options.get("file_url"))
    monthly = options.get("layout") == "Monthly"
    update = options.get("mode") == "Update Matching Rows"
    preserve_monthly = doc.get("entry_basis") == "Monthly" or doc.distribution_method == "Manual Monthly"
    normalize(doc)
    existing = {r.row_key: r for r in doc.proposal_details}
    months = {(m.row_key, m.month_number): m for m in doc.monthly_details}
    seen, errors = set(), []
    for idx, data in enumerate(rows, 2):
        try:
            for field in DIMENSIONS:
                data[field] = str(data.get(field) or "").strip()
            for field in ("sales_person", "territory", "customer_group"):
                data[field] = data[field] or doc.get(field) or ""
            if not data["item_code"]:
                raise ValueError("Item Code is required")
            key = make_row_key(*(data[f] for f in DIMENSIONS))
            n = number(data.get("month"), "Month") if monthly else 0
            if monthly and (n != int(n) or not 1 <= n <= 12):
                raise ValueError("Month must be 1–12")
            identity = (key, int(n))
            if identity in seen:
                raise ValueError("Duplicate target in import file")
            if not update and ((key, int(n)) in months if monthly else key in existing):
                raise ValueError("Target already exists; use Update Matching Rows")
            seen.add(identity)
            qty = number(data.get("target_qty"), "Target Qty")
            rate = number(data.get("target_rate"), "Target Rate")
            amount = number(data.get("target_amount"), "Target Amount") if data.get("target_amount") not in (None, "") else qty * rate
            if data.get("target_qty") in (None, "") or (data.get("target_rate") in (None, "") and data.get("target_amount") in (None, "")):
                raise ValueError("Target Qty and Target Rate or Target Amount are required")
            if data.get("target_rate") not in (None, "") and data.get("target_amount") not in (None, "") and abs(qty * rate - amount) > 0.01:
                raise ValueError("Quantity × Rate does not equal Amount")
            row = existing.get(key)
            if not row:
                row = doc.append("proposal_details", {**{f: data[f] for f in DIMENSIONS}, "row_key": key})
                existing[key] = row
            row.remarks = str(data.get("remarks") or "")
            row.override_reason = "Imported targets"
            if monthly:
                m = months.get(identity)
                if not m:
                    m = doc.append("monthly_details", {"row_key": key, "month_number": int(n)})
                    months[identity] = m
                m.target_qty, m.target_amount, m.target_rate = qty, amount, amount / qty if qty else 0
            else:
                row.approved_target_qty, row.approved_target_amount = qty, amount
                row.approved_target_rate = amount / qty if qty else 0
                if preserve_monthly:
                    from sales_performance.services.distribution_engine import equal_monthly
                    from sales_performance.services.precision import get_qty_precision, get_currency_precision
                    for allocation in equal_monthly(qty, amount, get_qty_precision(), get_currency_precision()):
                        month_key = (key, allocation["month_number"])
                        m = months.get(month_key)
                        if not m:
                            m = doc.append("monthly_details", {"row_key": key, "month_number": allocation["month_number"]})
                            months[month_key] = m
                        m.target_qty = allocation["target_qty"]
                        m.target_amount = allocation["target_amount"]
                        m.target_rate = allocation["target_rate"]
        except (ValueError, frappe.ValidationError) as exc:
            errors.append(f"Row {idx}: {exc}")
    if errors:
        frappe.throw("Import not applied:\n" + "\n".join(errors[:50]))
    doc.target_source, doc.value_basis = "Manual", "Quantity and Amount"
    doc.entry_basis = "Monthly" if monthly or preserve_monthly else "Annual"
    doc.distribution_method = "Manual Monthly" if monthly or preserve_monthly else "Equal Monthly"
    normalize(doc)
    return {"imported_rows": len(rows)}
